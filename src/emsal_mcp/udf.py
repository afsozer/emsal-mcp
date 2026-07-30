from __future__ import annotations

import html
import os
import re
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import Optional

from .models import build_error


class UdfError(ValueError):
    pass


UDF_TOOLKIT_VERSION = "0.9.0"
COMMAND_TIMEOUT_SECONDS = 30
DEFAULT_UDF_TOOLKIT_REPO = "https://github.com/saidsurucu/UDF-Toolkit.git"
BUNDLED_UDF_TOOLKIT_DIR = Path(__file__).resolve().parents[2] / "vendor" / "UDF-Toolkit"

UDF_AUTHORING_NOTE = (
    "UDF output follows the UYAP Dokuman Editor template structure (format_id 1.8)."
)


def _resolve_toolkit_dir() -> Optional[Path]:
    """Resolve UDF-Toolkit directory using local-yargi's candidate order."""
    candidates: list[str] = []
    for env_var in ("EMSAL_UDF_TOOLKIT_DIR", "UDF_TOOLKIT_PATH", "UDF_TOOLKIT_DIR"):
        val = os.environ.get(env_var, "").strip()
        if val:
            candidates.append(val)

    candidates.extend(
        [
            str(BUNDLED_UDF_TOOLKIT_DIR),
            str(Path.cwd() / "tools" / "UDF-Toolkit"),
            str(Path.home() / "tools" / "UDF-Toolkit"),
            str(Path.home() / "UDF-Toolkit"),
            r"C:\Users\Sozer\tools\UDF-Toolkit",
            r"C:\Users\Sozer\UDF-Toolkit",
        ]
    )

    for raw in candidates:
        p = Path(raw).expanduser().resolve()
        if p.is_dir() and (p / "udf_to_docx.py").exists():
            return p
    return None


def _looks_like_udf_toolkit(path: Path) -> bool:
    return path.is_dir() and (path / "udf_to_docx.py").exists()


def install_udf_toolkit(
    target_dir: str | Path | None = None,
    repo_url: str = DEFAULT_UDF_TOOLKIT_REPO,
    *,
    update: bool = False,
) -> dict:
    """Install or update the managed UDF-Toolkit clone."""
    target = Path(target_dir) if target_dir else BUNDLED_UDF_TOOLKIT_DIR
    target = target.expanduser().resolve()

    if _looks_like_udf_toolkit(target) and not update:
        return {
            "ok": True,
            "action": "install_udf_toolkit",
            "installed": False,
            "updated": False,
            "toolkit_dir": str(target),
            "message": "UDF-Toolkit is already installed.",
            "status": get_udf_toolkit_status(),
        }

    git_cmd = shutil.which("git")
    if not git_cmd:
        return build_error(
            "GIT_UNAVAILABLE",
            "git is required to install UDF-Toolkit automatically.",
            action="install_udf_toolkit",
            target_dir=str(target),
            repo_url=repo_url,
        )

    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        if target.exists():
            if update and (target / ".git").exists():
                result = subprocess.run(
                    [git_cmd, "-C", str(target), "pull", "--ff-only"],
                    capture_output=True,
                    text=True,
                    timeout=120,
                    check=False,
                )
                if result.returncode != 0:
                    return build_error(
                        "GIT_PULL_FAILED",
                        result.stderr.strip() or result.stdout.strip(),
                        action="install_udf_toolkit",
                        stdout=result.stdout.strip(),
                        stderr=result.stderr.strip(),
                        target_dir=str(target),
                    )
                return {
                    "ok": True,
                    "action": "install_udf_toolkit",
                    "installed": False,
                    "updated": True,
                    "toolkit_dir": str(target),
                    "stdout": result.stdout.strip(),
                    "status": get_udf_toolkit_status(),
                }
            return build_error(
                "TARGET_EXISTS_NOT_TOOLKIT",
                f"Target exists but does not look like UDF-Toolkit: {target}",
                action="install_udf_toolkit",
                target_dir=str(target),
            )

        result = subprocess.run(
            [git_cmd, "clone", "--depth", "1", repo_url, str(target)],
            capture_output=True,
            text=True,
            timeout=180,
            check=False,
        )
        if result.returncode != 0:
            return build_error(
                "GIT_CLONE_FAILED",
                result.stderr.strip() or result.stdout.strip(),
                action="install_udf_toolkit",
                stdout=result.stdout.strip(),
                stderr=result.stderr.strip(),
                target_dir=str(target),
                repo_url=repo_url,
            )
        if not _looks_like_udf_toolkit(target):
            return build_error(
                "INVALID_TOOLKIT_REPO",
                f"Cloned repository does not contain udf_to_docx.py: {target}",
                action="install_udf_toolkit",
                target_dir=str(target),
                repo_url=repo_url,
            )
        return {
            "ok": True,
            "action": "install_udf_toolkit",
            "installed": True,
            "updated": False,
            "toolkit_dir": str(target),
            "stdout": result.stdout.strip(),
            "status": get_udf_toolkit_status(),
        }
    except subprocess.TimeoutExpired:
        return build_error(
            "INSTALL_TIMEOUT",
            "Timed out while installing UDF-Toolkit.",
            action="install_udf_toolkit",
            target_dir=str(target),
            repo_url=repo_url,
        )
    except Exception as exc:
        return build_error(
            "INSTALL_FAILED",
            str(exc),
            action="install_udf_toolkit",
            target_dir=str(target),
            repo_url=repo_url,
        )


_discover_libreoffice = staticmethod(
    lambda: shutil.which("soffice") or shutil.which("libreoffice") or shutil.which("soffice.exe")
)
_discover_unoconv = staticmethod(lambda: shutil.which("unoconv"))
_discover_python = staticmethod(lambda: shutil.which("python") or shutil.which("python3") or shutil.which("py"))


def _run_text_command(args: list[str], timeout: int = 10) -> subprocess.CompletedProcess[str]:
    return subprocess.run(args, capture_output=True, text=True, timeout=timeout, check=False)


def get_udf_toolkit_status() -> dict:
    """Check whether external UDF-Toolkit scripts are available."""
    toolkit_dir = _resolve_toolkit_dir()
    env_enabled = os.environ.get("LOCAL_YARGI_UDF_TOOLS") == "1" or os.environ.get("EMSAL_UDF_TOOLS") == "1"
    disabled = os.environ.get("EMSAL_UDF_TOOLS") == "0" or os.environ.get("LOCAL_YARGI_UDF_TOOLS") == "0"
    managed = bool(toolkit_dir and toolkit_dir == BUNDLED_UDF_TOOLKIT_DIR.resolve())
    enabled = bool((env_enabled or managed) and not disabled)
    python_cmd = _discover_python()
    warnings: list[str] = []

    scripts_found = {
        "udf_to_docx": bool(toolkit_dir and (toolkit_dir / "udf_to_docx.py").exists()),
        "udf_to_pdf": bool(toolkit_dir and (toolkit_dir / "udf_to_pdf.py").exists()),
        "docx_to_udf": bool(toolkit_dir and (toolkit_dir / "docx_to_udf.py").exists()),
    }
    dependencies = {
        "python-docx": False,
        "PyMuPDF": False,
        "Pillow": False,
        "reportlab": False,
    }

    python_version: str | None = None
    if python_cmd:
        try:
            version_result = _run_text_command([python_cmd, "--version"], timeout=5)
            python_version = (version_result.stdout or version_result.stderr).strip() or None
        except Exception as exc:
            warnings.append(f"Python version check failed: {exc}")

        dep_code = (
            "deps = ['docx', 'fitz', 'PIL', 'reportlab']; out = []\n"
            "for dep in deps:\n"
            "    try:\n"
            "        __import__(dep); out.append(dep + ':1')\n"
            "    except ImportError:\n"
            "        out.append(dep + ':0')\n"
            "print(','.join(out))"
        )
        try:
            dep_result = _run_text_command([python_cmd, "-c", dep_code], timeout=10)
            for part in dep_result.stdout.strip().split(","):
                name, _, val = part.partition(":")
                ok = val == "1"
                if name == "docx":
                    dependencies["python-docx"] = ok
                elif name == "fitz":
                    dependencies["PyMuPDF"] = ok
                elif name == "PIL":
                    dependencies["Pillow"] = ok
                elif name == "reportlab":
                    dependencies["reportlab"] = ok
        except Exception as exc:
            warnings.append(f"Python dependency check failed: {exc}")
    else:
        warnings.append("Python ('python', 'python3' or 'py') was not found on PATH.")

    if not toolkit_dir:
        warnings.append("UDF-Toolkit was not found or toolkit path is not configured.")
    elif not enabled:
        warnings.append(
            f"UDF-Toolkit detected at {toolkit_dir}, but integration is disabled. "
            "Set LOCAL_YARGI_UDF_TOOLS=1 or EMSAL_UDF_TOOLS=1, or remove EMSAL_UDF_TOOLS=0."
        )
    else:
        if not scripts_found["udf_to_docx"]:
            warnings.append("udf_to_docx.py was not found under UDF-Toolkit.")
        if not scripts_found["udf_to_pdf"]:
            warnings.append("udf_to_pdf.py was not found under UDF-Toolkit.")
        if not scripts_found["docx_to_udf"]:
            warnings.append("docx_to_udf.py was not found under UDF-Toolkit.")

    if python_cmd:
        missing = [name for name, ok in dependencies.items() if not ok]
        if missing:
            warnings.append(f"Missing Python packages: {', '.join(missing)}.")

    configured = bool(enabled and toolkit_dir)
    ok = bool(configured and python_cmd and all(scripts_found.values()))

    return {
        "ok": ok,
        "enabled": enabled,
        "configured": configured,
        "pythonAvailable": python_cmd is not None,
        "pythonVersion": python_version,
        "pythonCmd": python_cmd,
        "toolkitPath": str(toolkit_dir) if toolkit_dir else None,
        "toolkit_dir": str(toolkit_dir) if toolkit_dir else None,
        "managedToolkitPath": str(BUNDLED_UDF_TOOLKIT_DIR.resolve()),
        "managed": managed,
        "scriptsFound": scripts_found,
        "dependencies": dependencies,
        "libreoffice_path": _discover_libreoffice(),
        "unoconv_path": _discover_unoconv(),
        "version": UDF_TOOLKIT_VERSION,
        "warnings": warnings,
    }


def get_udf_authoring_instructions(format: str = "json") -> dict:
    steps = [
        "For DOCX sources, convert directly (docx_to_udf_native / export_document with docx_path) to preserve formatting.",
        "For plain text, write_udf applies Turkish petition formatting automatically.",
        "Keep placeholders in {{...}} format.",
    ]
    warnings = [
        UDF_AUTHORING_NOTE,
    ]

    if format == "markdown":
        lines = ["# UDF Authoring Instructions", "", "## Steps"]
        lines.extend(f"{idx}. {step}" for idx, step in enumerate(steps, 1))
        lines.extend(["", "## Warnings"])
        lines.extend(f"- {warning}" for warning in warnings)
        return {
            "format": "markdown",
            "instructions": "\n".join(lines),
            "warnings": warnings,
            "version": UDF_TOOLKIT_VERSION,
        }

    return {
        "format": "json",
        "steps": steps,
        "warnings": warnings,
        "version": UDF_TOOLKIT_VERSION,
        "toolkit_status": get_udf_toolkit_status(),
    }


def read_udf(path: str | Path) -> str:
    """Read and extract text content from a UYAP UDF file."""
    p = Path(path)
    if not p.exists():
        raise UdfError(f"UDF bulunamadi: {p}")
    if p.read_bytes()[:2] != b"PK":
        raise UdfError("Gecersiz UDF: ZIP imzasi yok")

    try:
        with zipfile.ZipFile(p) as zf:
            if "content.xml" not in zf.namelist():
                raise UdfError("Invalid UDF: content.xml missing")
            xml = zf.read("content.xml").decode("utf-8", errors="ignore")
    except UdfError:
        raise
    except Exception as exc:
        raise UdfError(f"Failed to parse UDF ZIP structure: {exc}") from exc

    cdata = re.findall(r"<content>\s*<!\[CDATA\[([\s\S]*?)\]\]>\s*</content>", xml)
    if cdata:
        raw_text = "\n".join(cdata)
    else:
        contents = re.findall(r"<content>([\s\S]*?)</content>", xml)
        if contents:
            raw_text = "\n".join(contents)
        else:
            paragraphs = re.findall(r"<paragraph[^>]*>([\s\S]*?)</paragraph>", xml)
            raw_text = "\n".join(paragraphs) if paragraphs else xml

    return html.unescape(re.sub(r"<[^>]*>", "", raw_text)).strip()


def udf_to_markdown(path: str | Path) -> str:
    """Convert a UDF file content to basic markdown paragraphs."""
    lines = [ln.strip() for ln in read_udf(path).splitlines()]
    paragraphs: list[str] = []
    cur: list[str] = []
    for line in lines:
        if line:
            cur.append(line)
        elif cur:
            paragraphs.append(" ".join(cur))
            cur = []
    if cur:
        paragraphs.append(" ".join(cur))
    return "\n\n".join(paragraphs)


# Formatting model: each paragraph is a dict with keys:
#   align:    UDF Alignment code ("0" left, "1" center, "2" right, "3" justify)
#   runs:     list of (text, bold, italic, underline) covering the text exactly
#   indent:   LeftIndent in points (0.0 = none)
#   spacing:  UDF LineSpacing (line spacing minus 1.0)
#   numbered: None, or (list_id, level) for UYAP numbered list paragraphs
# The element layout mirrors UDF files produced by UYAP Dokuman Editor itself:
# paragraphs joined with "\n" in the CDATA pool, the newline carried by the
# paragraph's last run.

DEFAULT_LINE_SPACING = 0.14999998

_TR_UPPER = "A-ZÇĞİÖŞÜ"
_LABEL_RE = re.compile(rf"^([{_TR_UPPER}][{_TR_UPPER}0-9 ./()&-]*?\t+)(:.*)$")
_ENUM_RE = re.compile(r"^(\d{1,2}[-.)]\s+|[a-zçğıöşü][.)]\s+)(.*)$")
_CAPS_LINE_RE = re.compile(rf"^[{_TR_UPPER}0-9IVXLC ().,:;'’\"/&\t–—-]+$")
# Caption labels up to 9 chars ("DOSYA NO", "VEKİLLERİ", ...) followed by a
# single tab and colon; UYAP's wide default tab stops need a second tab for
# the colons to line up. Longer labels ("SONUÇ VE TALEP") already reach the
# stop with one tab.
_SHORT_LABEL_TAB_RE = re.compile(rf"^[{_TR_UPPER}][{_TR_UPPER}0-9 ./()&-]{{0,8}}\t:")

# Blank spacer paragraphs are inserted where the DOCX relied on space-after /
# space-before (UDF has no such concept). 5pt catches the petition body
# rhythm (5-6pt after paragraphs, 8-12pt around headings) while leaving the
# caption block (3pt), numbered request lists (4pt) and EK lists (2pt) tight.
SPACER_GAP_PT = 5.0


def _para_style(
    align: str = "3",
    runs: list[tuple[str, bool, bool, bool]] | None = None,
    indent: float = 0.0,
    spacing: float = DEFAULT_LINE_SPACING,
    numbered: tuple[int, int] | None = None,
) -> dict:
    return {
        "align": align,
        "runs": runs if runs is not None else [("", False, False, False)],
        "indent": indent,
        "spacing": spacing,
        "numbered": numbered,
    }


def _is_heading_line(line: str) -> bool:
    stripped = line.strip()
    if len(stripped) < 3 or not re.search(rf"[{_TR_UPPER}]", stripped):
        return False
    return bool(_CAPS_LINE_RE.match(stripped))


def _smart_paragraphs(paragraphs: list[str], title_centered: bool) -> list[dict]:
    """Infer dilekçe-style formatting for plain-text paragraphs.

    Returns a list of paragraph-style dicts matching `paragraphs`.
    Conventions (modelled on UYAP-authored petitions):
    - first non-empty line: centered title, bold
    - "ETİKET<tab>: değer" lines: label bold, value plain
    - ALL-CAPS lines (section headings): bold
    - "1-" / "a)" enumerators: enumerator bold
    - trailing short signature lines containing "Vekili" / "Av.": right, bold
    """
    styled: list[dict] = []
    first_text_idx = next((i for i, p in enumerate(paragraphs) if p.strip()), -1)

    # Detect a trailing signature block: up to 4 short trailing lines where at
    # least one contains "Vekili" or starts with "Av.".
    signature_idx: set[int] = set()
    tail: list[int] = []
    for i in range(len(paragraphs) - 1, -1, -1):
        if not paragraphs[i].strip():
            if tail:
                break
            continue
        if len(paragraphs[i].strip()) > 45 or len(tail) >= 4:
            break
        tail.append(i)
    if any("Vekili" in paragraphs[i] or paragraphs[i].strip().startswith("Av.") for i in tail):
        signature_idx = set(tail)

    for idx, para in enumerate(paragraphs):
        align = "3"
        runs: list[tuple[str, bool, bool, bool]] = []
        if not para:
            runs = [("", False, False, False)]
        elif idx == first_text_idx and title_centered:
            align = "1"
            runs = [(para, True, False, False)]
        elif idx in signature_idx:
            align = "2"
            runs = [(para, True, False, False)]
        elif _LABEL_RE.match(para):
            m = _LABEL_RE.match(para)
            runs = [(m.group(1), True, False, False), (m.group(2), False, False, False)]
        elif _is_heading_line(para):
            runs = [(para, True, False, False)]
        else:
            m = _ENUM_RE.match(para)
            if m:
                runs = [(m.group(1), True, False, False), (m.group(2), False, False, False)]
            else:
                runs = [(para, False, False, False)]
        styled.append(_para_style(align=align, runs=runs))
    return styled


# Page-number footer as emitted by UYAP Dokuman Editor (bold Times New Roman
# 12, "current/total" style). The footer paragraph references a 3-char "  \n"
# tail appended to the CDATA pool.
_FOOTER_XML_TMPL = (
    '<footer pageNumber-spec="BSP32_2120" pageNumber-seperator="/" pageNumber-fontBold="true" '
    'pageNumber-fontItalic="false" pageNumber-fontFace="Times New Roman" pageNumber-fontSize="12" '
    'pageNumber-color="-16777216" color-boundary="-1" pageNumber-foreStr="" pageNumber-pageStartNumStr="">'
    '<paragraph name="hvl-default" family="Times New Roman" size="12" description="Gövde">'
    '<content name="hvl-default" family="Times New Roman" size="12" description="Gövde" '
    'startOffset="{offset}" length="3" /></paragraph></footer>'
)


def _render_udf_xml(
    paragraphs: list[str],
    styled: list[dict],
    *,
    page_number_footer: bool = False,
) -> str:
    pool = "\n".join(paragraphs)
    elements: list[str] = []
    offset = 0
    last = len(paragraphs) - 1
    for idx, (para, st) in enumerate(zip(paragraphs, styled)):
        attrs = f'Alignment="{st["align"]}" LineSpacing="{st["spacing"]}"'
        if st["numbered"]:
            list_id, level = st["numbered"]
            attrs += (
                ' Numbered="true" NumberType="NUMBER_TYPE_NUMBER_DOT"'
                f' ListLevel="{level}" ListId="{list_id}"'
            )
        if st["indent"]:
            attrs += f' LeftIndent="{st["indent"]}"'
        parts: list[str] = []
        # Trailing newline (paragraph separator) rides on the last run. With a
        # footer the last visible paragraph also carries one, since the pool
        # continues with the footer text.
        trailing = 1 if (idx < last or page_number_footer) else 0
        runs = st["runs"]
        nonempty = [r for r in runs if r[0]] or list(runs[-1:]) or [("", False, False, False)]
        for ridx, (rtext, bold, italic, underline) in enumerate(nonempty):
            length = len(rtext) + (trailing if ridx == len(nonempty) - 1 else 0)
            if length <= 0:
                continue
            rattrs = ""
            if bold:
                rattrs += ' bold="true"'
            if italic:
                rattrs += ' italic="true"'
            if underline:
                rattrs += ' underline="true"'
            parts.append(f'<content{rattrs} startOffset="{offset}" length="{length}" />')
            offset += length
        if not parts:  # empty final paragraph
            parts.append(f'<content startOffset="{offset}" length="0" />')
        elements.append(f"<paragraph {attrs}>{''.join(parts)}</paragraph>")

    if page_number_footer:
        # Pool tail "\n\n  \n": paragraph separator (owned by the last run),
        # one filler newline, then the "  \n" the footer paragraph points at —
        # byte-for-byte what UYAP Dokuman Editor produces.
        footer_offset = len(pool) + 2
        pool += "\n\n  \n"
        elements.append(_FOOTER_XML_TMPL.format(offset=footer_offset))

    return f'''<?xml version="1.0" encoding="UTF-8" ?>
<template format_id="1.8">
<content><![CDATA[{pool}]]></content>
<properties><pageFormat mediaSizeName="1" leftMargin="56.69" rightMargin="56.69" topMargin="56.69" bottomMargin="56.69" paperOrientation="1" headerFOffset="20.0" footerFOffset="20.0" /></properties>
<elements resolver="hvl-default">
{''.join(elements)}
</elements>
<styles><style name="default" description="Geçerli" family="Dialog" size="12" bold="false" italic="false" foreground="-13421773" FONT_ATTRIBUTE_KEY="javax.swing.plaf.FontUIResource[family=Dialog,name=Dialog,style=plain,size=12]" /><style name="hvl-default" family="Times New Roman" size="12" description="Gövde" /></styles>
</template>'''


def write_udf(
    text: str,
    out_path: str | Path,
    *,
    title_centered: bool = False,
    smart: bool = True,
) -> Path:
    """Write a simple UDF package. Caller must verify in UYAP Dokuman Editor.

    With smart=True (default), Turkish petition conventions are applied:
    bold centered title, bold "ETİKET\t:" labels, bold ALL-CAPS headings,
    bold enumerators and a right-aligned bold signature block.
    """
    out = Path(out_path)
    paragraphs = text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    if smart:
        styled = _smart_paragraphs(paragraphs, title_centered)
    else:
        styled = [
            _para_style(
                align="1" if idx == 0 and title_centered else "3",
                runs=[(para, idx == 0 and title_centered, False, False)],
            )
            for idx, para in enumerate(paragraphs)
        ]
    xml = _render_udf_xml(paragraphs, styled)
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("content.xml", xml.encode("utf-8"))
    return out


def _docx_numbering(para) -> tuple[int, int] | None:
    """Extract (list_id, level) from a DOCX paragraph's w:numPr, if any."""
    from docx.oxml.ns import qn

    pPr = para._p.pPr
    if pPr is None:
        return None
    numPr = pPr.find(qn("w:numPr"))
    if numPr is None:
        return None

    def _val(tag: str, default: int) -> int:
        el = numPr.find(qn(tag))
        try:
            return int(el.get(qn("w:val"))) if el is not None else default
        except (TypeError, ValueError):
            return default

    num_id = _val("w:numId", 0)
    if num_id <= 0:  # numId 0 means "numbering removed"
        return None
    return (num_id, _val("w:ilvl", 0) + 1)


def _iter_docx_runs(para):
    """Yield a paragraph's runs including those inside w:hyperlink elements.

    python-docx's paragraph.runs skips hyperlink content entirely, which
    silently drops link text (e.g. "link: https://..." in EK lists).
    """
    if hasattr(para, "iter_inner_content"):
        for item in para.iter_inner_content():
            if hasattr(item, "runs"):  # Hyperlink
                yield from item.runs
            else:
                yield item
    else:  # older python-docx without iter_inner_content
        yield from para.runs


def _double_short_label_tab(
    runs: list[tuple[str, bool, bool, bool]], text: str
) -> tuple[list[tuple[str, bool, bool, bool]], str]:
    """Turn "ETİKET\\t:" into "ETİKET\\t\\t:" for short caption labels."""
    if text.count("\t") != 1 or not _SHORT_LABEL_TAB_RE.match(text):
        return runs, text
    new_runs = [
        (rtext.replace("\t", "\t\t", 1) if "\t" in rtext else rtext, b, i, u)
        for rtext, b, i, u in runs
    ]
    return new_runs, text.replace("\t", "\t\t", 1)


def docx_to_udf_native(
    docx_path: str | Path,
    out_path: str | Path | None = None,
    *,
    spacer_paragraphs: bool = True,
    label_double_tab: bool = True,
    page_number_footer: bool = True,
) -> dict:
    """Convert DOCX to UDF preserving formatting, without external tools.

    Carries over per-run bold, italic and underline, paragraph alignment,
    left indent, line spacing and numbered lists (w:numPr -> UYAP Numbered
    paragraphs). Element layout mirrors UYAP Dokuman Editor output
    (newline-joined CDATA pool). Tables and images are not supported.

    UDF has no space-after/space-before concept, so with spacer_paragraphs
    (default) an empty paragraph is inserted wherever the DOCX vertical gap is
    >= SPACER_GAP_PT — except around numbered list items and before
    right-aligned (signature) paragraphs, which stay tight. label_double_tab
    aligns caption label colons on UYAP's tab stops, and page_number_footer
    adds the standard "page/total" footer.
    """
    action = "docx_to_udf_native"
    src = Path(docx_path)
    if not src.exists():
        return build_error("FILE_NOT_FOUND", f"DOCX file not found: {src}", action=action)
    try:
        import docx as _docx  # python-docx
        from docx.enum.text import WD_ALIGN_PARAGRAPH as _WD
    except ImportError:
        return build_error(
            "PYTHON_DOCX_MISSING",
            "python-docx is required for native DOCX -> UDF conversion.",
            action=action,
        )

    try:
        document = _docx.Document(str(src))
    except Exception as exc:
        return build_error("DOCX_READ_FAILED", str(exc), action=action)

    align_map = {
        _WD.LEFT: "0",
        _WD.CENTER: "1",
        _WD.RIGHT: "2",
        _WD.JUSTIFY: "3",
    }

    paragraphs: list[str] = []
    styled: list[dict] = []
    gaps_after: list[float] = []
    gaps_before: list[float] = []
    for para in document.paragraphs:
        pf = para.paragraph_format
        align = align_map.get(para.alignment, "3")
        numbered = _docx_numbering(para)
        indent = round(pf.left_indent.pt, 2) if pf.left_indent else 0.0
        if numbered and not indent:
            indent = 25.0  # UYAP's default list indent
        spacing = DEFAULT_LINE_SPACING
        try:
            ls = pf.line_spacing
            if isinstance(ls, float) and ls > 1.0:
                spacing = round(ls - 1.0, 8)
        except Exception:
            pass
        gaps_after.append(pf.space_after.pt if pf.space_after is not None else 0.0)
        gaps_before.append(pf.space_before.pt if pf.space_before is not None else 0.0)

        runs: list[tuple[str, bool, bool, bool]] = []
        for run in _iter_docx_runs(para):
            rtext = run.text
            if not rtext:
                continue
            bold = bool(run.bold if run.bold is not None else run.font.bold)
            italic = bool(run.italic if run.italic is not None else run.font.italic)
            underline = bool(run.underline if run.underline is not None else run.font.underline)
            if runs and runs[-1][1:] == (bold, italic, underline):
                runs[-1] = (runs[-1][0] + rtext, bold, italic, underline)
            else:
                runs.append((rtext, bold, italic, underline))
        text = "".join(r[0] for r in runs)
        if not runs:
            runs = [("", False, False, False)]
        if label_double_tab:
            runs, text = _double_short_label_tab(runs, text)
        paragraphs.append(text)
        styled.append(
            _para_style(align=align, runs=runs, indent=indent, spacing=spacing, numbered=numbered)
        )

    if not paragraphs:
        return build_error("DOCX_EMPTY", "DOCX contains no paragraphs.", action=action)

    spacers_added = 0
    if spacer_paragraphs:
        merged_paragraphs: list[str] = []
        merged_styled: list[dict] = []
        for idx, (text, st) in enumerate(zip(paragraphs, styled)):
            merged_paragraphs.append(text)
            merged_styled.append(st)
            if idx >= len(paragraphs) - 1:
                continue
            gap = gaps_after[idx] + gaps_before[idx + 1]
            nxt = styled[idx + 1]
            if (
                gap >= SPACER_GAP_PT
                and text.strip()
                and paragraphs[idx + 1].strip()
                and not st["numbered"]
                and not nxt["numbered"]
                and nxt["align"] != "2"  # keep signature blocks attached
            ):
                _, lb, li, lu = st["runs"][-1]
                merged_paragraphs.append("")
                merged_styled.append(
                    _para_style(
                        align=st["align"],
                        runs=[("", lb, li, lu)],
                        indent=st["indent"],
                        spacing=st["spacing"],
                    )
                )
                spacers_added += 1
        paragraphs, styled = merged_paragraphs, merged_styled

    out = Path(out_path) if out_path else src.with_suffix(".udf")
    xml = _render_udf_xml(paragraphs, styled, page_number_footer=page_number_footer)
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("content.xml", xml.encode("utf-8"))
    return {
        "ok": True,
        "action": action,
        "out_path": str(out),
        "file_size": out.stat().st_size,
        "paragraphs": len(paragraphs),
        "spacers_added": spacers_added,
        "page_number_footer": page_number_footer,
    }


def probe_udf(path: str | Path) -> dict:
    """Probe a UDF file using the local-yargi-style status shape."""
    p = Path(path)
    absolute = p.resolve()
    info: dict = {
        "filePath": str(absolute),
        "path": str(p),
        "exists": p.exists(),
        "isUdfLike": p.suffix.lower() == ".udf",
        "isZipLike": False,
        "zip": False,
        "hasContentXml": False,
        "has_content_xml": False,
        "contentXmlSize": 0,
        "parseableXml": False,
        "textCharCount": 0,
        "text_length": 0,
        "format_id": None,
        "content_xml_preview": None,
        "warnings": [],
        "recommendedNextStep": "",
    }

    if not p.exists():
        info["warnings"].append("File does not exist on the local filesystem.")
        info["recommendedNextStep"] = "Check the file path and try again."
        return info

    data = p.read_bytes()
    if not data:
        info["warnings"].append("File is empty (0 bytes).")
        info["recommendedNextStep"] = "Ensure the file contains valid UDF data."
        return info

    info["isZipLike"] = len(data) >= 4 and data[:4] == b"PK\x03\x04"
    info["zip"] = info["isZipLike"]
    if not info["isZipLike"]:
        info["warnings"].append("File lacks standard ZIP local header signature (PK\\x03\\x04).")

    try:
        with zipfile.ZipFile(p) as zf:
            names = zf.namelist()
            info["hasContentXml"] = "content.xml" in names
            info["has_content_xml"] = info["hasContentXml"]
            if not info["hasContentXml"]:
                info["warnings"].append("content.xml was not found inside the ZIP package.")
                info["recommendedNextStep"] = "Make sure this is a genuine UYAP UDF document."
                return info

            zinfo = zf.getinfo("content.xml")
            info["contentXmlSize"] = zinfo.file_size
            xml = zf.read("content.xml").decode("utf-8", errors="ignore")
            info["content_xml_preview"] = xml[:500]
            m = re.search(r'<template[^>]+format_id=["\']([^"\']+)', xml)
            info["format_id"] = m.group(1) if m else None

        text = read_udf(p)
        info["parseableXml"] = True
        info["textCharCount"] = len(text)
        info["text_length"] = len(text)
        if text:
            info["recommendedNextStep"] = "Use 'emsal-mcp udf read' or 'emsal-mcp udf to-md' to read document content."
        else:
            info["warnings"].append("UDF content extracted successfully but contains no readable text.")
    except Exception as exc:
        info["warnings"].append(str(exc))
        info["recommendedNextStep"] = "Verify file integrity or check with UYAP Dokuman Editor."

    return info


def _toolkit_unavailable(action: str) -> dict:
    status = get_udf_toolkit_status()
    return build_error(
        "TOOLKIT_UNAVAILABLE",
        "UDF-Toolkit is not enabled or properly configured.",
        action=action,
        toolkit_status=status,
        recommended_next_steps=[
            "Set LOCAL_YARGI_UDF_TOOLS=1 or EMSAL_UDF_TOOLS=1.",
            "Set UDF_TOOLKIT_PATH or EMSAL_UDF_TOOLKIT_DIR to a UDF-Toolkit clone.",
        ],
    )


def _ensure_udf_toolkit_for_conversion() -> dict:
    status = get_udf_toolkit_status()
    if status["enabled"] and status["configured"]:
        return status
    if os.environ.get("EMSAL_UDF_AUTO_INSTALL") == "0":
        return status

    installed = install_udf_toolkit()
    if not installed.get("ok"):
        status["autoInstall"] = installed
        return status
    return get_udf_toolkit_status()


def _run_udf_toolkit_script(script_name: str, args: list[str], action: str) -> dict:
    status = _ensure_udf_toolkit_for_conversion()
    if not status["enabled"] or not status["configured"]:
        return _toolkit_unavailable(action)
    if not status["pythonAvailable"] or not status["pythonCmd"]:
        return build_error(
            "PYTHON_UNAVAILABLE",
            "Python is not available to run UDF-Toolkit.",
            action=action,
            toolkit_status=status,
        )

    toolkit_dir = Path(status["toolkitPath"])
    script_path = toolkit_dir / script_name
    if not script_path.exists():
        return build_error(
            "SCRIPT_NOT_FOUND",
            f"Script not found: {script_path}",
            action=action,
            toolkit_status=status,
        )

    try:
        result = subprocess.run(
            [status["pythonCmd"], str(script_path), *args],
            cwd=toolkit_dir,
            capture_output=True,
            text=True,
            timeout=COMMAND_TIMEOUT_SECONDS,
            check=False,
        )
    except subprocess.TimeoutExpired:
        return build_error(
            "TOOLKIT_TIMEOUT",
            f"UDF-Toolkit command timed out after {COMMAND_TIMEOUT_SECONDS}s.",
            action=action,
            toolkit_status=status,
        )
    except Exception as exc:
        return build_error("EXCEPTION", str(exc), action=action, toolkit_status=status)

    if result.returncode != 0:
        return build_error(
            "CONVERSION_FAILED",
            result.stderr.strip() or result.stdout.strip() or "UDF-Toolkit conversion failed",
            action=action,
            stdout=result.stdout.strip(),
            stderr=result.stderr.strip(),
            toolkit_status=status,
        )

    return {
        "ok": True,
        "action": action,
        "stdout": result.stdout.strip(),
        "stderr": result.stderr.strip(),
        "toolkit_status": status,
    }


def convert_udf_to_docx(file_path: str | Path, out_path: str | Path | None = None) -> dict:
    """Convert UDF to DOCX via UDF-Toolkit udf_to_docx.py."""
    action = "convert_udf_to_docx"
    src = Path(file_path)
    if not src.exists():
        return build_error("FILE_NOT_FOUND", f"UDF file not found: {src}", action=action)

    out = Path(out_path) if out_path else src.with_suffix(".docx")
    result = _run_udf_toolkit_script("udf_to_docx.py", [str(src), str(out)], action)
    if not result.get("ok"):
        return result
    return {
        **result,
        "out_path": str(out),
        "file_size": out.stat().st_size if out.exists() else 0,
    }


def convert_udf_to_pdf(file_path: str | Path, out_path: str | Path | None = None) -> dict:
    """Convert UDF to PDF via UDF-Toolkit udf_to_pdf.py."""
    action = "convert_udf_to_pdf"
    src = Path(file_path)
    if not src.exists():
        return build_error("FILE_NOT_FOUND", f"UDF file not found: {src}", action=action)

    out = Path(out_path) if out_path else src.with_suffix(".pdf")
    result = _run_udf_toolkit_script("udf_to_pdf.py", [str(src), str(out)], action)
    if not result.get("ok"):
        return result
    return {**result, "out_path": str(out), "file_size": out.stat().st_size if out.exists() else 0}


def convert_docx_to_udf(
    file_path: str | Path,
    out_path: str | Path | None = None,
) -> dict:
    """Convert DOCX to UDF, preserving formatting.

    Uses the native converter first (no external tools needed); falls back to
    the UDF-Toolkit docx_to_udf.py script if the native conversion fails.
    """
    action = "convert_docx_to_udf"
    src = Path(file_path)
    if not src.exists():
        return build_error("FILE_NOT_FOUND", f"DOCX file not found: {src}", action=action)

    out = Path(out_path) if out_path else src.with_suffix(".udf")

    native = docx_to_udf_native(src, out)
    if native.get("ok"):
        return {**native, "action": action, "converter": "native"}

    result = _run_udf_toolkit_script("docx_to_udf.py", [str(src), str(out)], action)
    if not result.get("ok"):
        return {**native, "action": action, "toolkit_fallback": result}

    if not out.exists():
        auto_out = src.with_suffix(".udf")
        if auto_out.exists():
            try:
                out.parent.mkdir(parents=True, exist_ok=True)
                auto_out.replace(out)
            except Exception as exc:
                return build_error(
                    "OUTPUT_MOVE_FAILED",
                    str(exc),
                    action=action,
                    stdout=result.get("stdout", ""),
                )
        else:
            return build_error(
                "OUTPUT_MISSING",
                f"Expected UDF output was not created: {out}",
                action=action,
                stdout=result.get("stdout", ""),
            )

    return {
        **result,
        "out_path": str(out),
        "file_size": out.stat().st_size if out.exists() else 0,
        "converter": "udf-toolkit",
    }


def convert_docx_to_udf_experimental(
    file_path: str | Path,
    out_path: str | Path | None = None,
    experimental: bool = False,
) -> dict:
    """Deprecated alias for convert_docx_to_udf; the flag is ignored."""
    return convert_docx_to_udf(file_path, out_path)
