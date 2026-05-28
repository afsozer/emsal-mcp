from __future__ import annotations

import json
from pathlib import Path
from typing import Iterable

from docx import Document as DocxDocument

from .models import Document
from .safety import build_input_pack, citation_check


def controlled_draft(title: str, body: str, docs: Iterable[Document]) -> str:
    safe_docs = [d for d in docs if citation_check(d).ok]
    lines = [f"# {title}", "", body.strip(), "", "## Güvenli Atıf Adayları"]
    if not safe_docs:
        lines.append("Güvenli tam metinli karar bulunamadı; künye/atıf uydurulmadı.")
    for doc in safe_docs:
        lines.append(f"- {doc.citation_label()} (`{doc.source}:{doc.document_id}`)")
    lines.append("\n> Taslak avukat denetimi gerektirir; eksik metadata tamamlanmadan atıf yapılmaz.")
    return "\n".join(lines)


def markdown_to_docx(markdown: str, out_path: str | Path) -> Path:
    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    doc = DocxDocument()
    for line in markdown.splitlines():
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:].strip(), style="List Bullet")
        elif line.strip():
            doc.add_paragraph(line)
        else:
            doc.add_paragraph("")
    doc.save(out)
    return out


def export_bundle(out_dir: str | Path, *, matter: str, issue: str, docs: list[Document]) -> dict:
    out = Path(out_dir)
    out.mkdir(parents=True, exist_ok=True)
    pack = build_input_pack(matter, issue, docs)
    (out / "input-pack.json").write_text(pack.model_dump_json(indent=2), encoding="utf-8")
    (out / "safe-citations.md").write_text("\n".join(f"- {c}" for c in pack.safe_citations), encoding="utf-8")
    (out / "excluded.json").write_text(json.dumps(pack.excluded, ensure_ascii=False, indent=2), encoding="utf-8")
    for doc in docs:
        if doc.markdown:
            name = f"{doc.source}-{doc.document_id}".replace("/", "_").replace(":", "_")[:120] + ".md"
            (out / name).write_text(doc.markdown, encoding="utf-8")
    return {"out_dir": str(out), "safe_count": len(pack.documents), "excluded_count": len(pack.excluded)}
